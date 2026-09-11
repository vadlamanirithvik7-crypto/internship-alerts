"""Persist the application workbook in the same transaction as progress changes."""

from poller.application_links import notification_url

import base64
from io import BytesIO
from sqlalchemy import select, or_
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from shared.db import AICache, Posting, utcnow

KEY = "application-workbook-v1"


def update_workbook(db):
    rows = db.scalars(
        select(Posting)
        .where(
            or_(
                Posting.applied_at.is_not(None),
                Posting.status.in_(["applied", "interview", "offer", "rejected"]),
            )
        )
        .order_by(Posting.applied_at.desc(), Posting.id.desc())
    ).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Applications"
    sheet.append(
        [
            "Posting ID",
            "Company",
            "Role",
            "Location",
            "Term",
            "Status",
            "Applied at (UTC)",
            "Updated at (UTC)",
            "Application URL",
            "Notes",
        ]
    )
    for p in rows:
        values = [
            p.id,
            p.company_name,
            p.title,
            p.location or "",
            p.term or "",
            p.status,
            p.applied_at.isoformat(sep=" ", timespec="seconds") if p.applied_at else "",
            p.status_updated_at.isoformat(sep=" ", timespec="seconds")
            if p.status_updated_at
            else "",
            notification_url(p),
            p.notes or "",
        ]
        sheet.append(values)
        # Source text and personal notes must never become Excel formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173D36")
    for column, width in zip("ABCDEFGHIJ", [12, 28, 48, 28, 22, 18, 24, 24, 70, 55]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    artifact = db.get(AICache, KEY)
    if artifact is None:
        artifact = AICache(key=KEY, kind="spreadsheet", payload="")
        db.add(artifact)
    artifact.payload = base64.b64encode(output.getvalue()).decode("ascii")
    artifact.created_at = utcnow()
    db.flush()
    return output.getvalue()


def workbook_bytes(db):
    artifact = db.get(AICache, KEY)
    if artifact:
        return base64.b64decode(artifact.payload)
    content = update_workbook(db)
    db.commit()
    return content
